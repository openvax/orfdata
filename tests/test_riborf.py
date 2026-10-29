"""Native mouse RibORF identity precedes exact target-RNA placement."""

import gzip
import hashlib
import json
from dataclasses import replace
from pathlib import Path

import pytest

from orfdata import (
    MS,
    CatalogManifest,
    CatalogRegistry,
    Proteome,
    RiboSeq,
    extra_orfs,
    load_catalog,
)
from orfdata.catalog import _record
from orfdata.sources.riborf_native import _coding, _fasta, reconstruct

FIXTURES = Path(__file__).parent / "data/riborf"
FILES = dict(
    table="SD3.xlsx",
    gtf="ensembl107.gtf.gz",
    cdna="ensembl107.cdna.fa.gz",
    ncrna="ensembl107.ncrna.fa.gz",
    source_gtf="gencodeM25.gtf.gz",
    source_rna="gencodeM25.rna.fa.gz",
)
PATHS = {key: FIXTURES / name for key, name in FILES.items()}


@pytest.fixture
def study(tmp_path):
    paths = {}
    for key, source in PATHS.items():
        paths[key] = tmp_path / source.name
        paths[key].write_bytes(source.read_bytes())
    manifest = CatalogManifest.from_dict(json.loads((FIXTURES / "v3-manifest.json").read_text()))
    return paths, manifest


def repin(paths, manifest):
    sources = tuple(
        replace(
            s,
            sha256=hashlib.sha256(paths[s.asset_id].read_bytes()).hexdigest(),
            size=paths[s.asset_id].stat().st_size,
        )
        for s in manifest.sources
    )
    hashes = {s.asset_id: s.sha256 for s in sources}
    return replace(
        manifest,
        sources=sources,
        source_references=(
            replace(
                manifest.source_references[0],
                sequence_sha256=hashes["source_rna"],
                source_sha256=hashes["source_gtf"],
            ),
        ),
    )


@pytest.fixture
def catalog(study, tmp_path):
    paths, manifest = study
    release = dict(
        manifest=manifest.to_dict(),
        distribution=dict(access_terms="fixture", redistributable=False),
        assets={},
    )
    registry = CatalogRegistry(
        {
            manifest.catalog_id: dict(
                default_version=manifest.version, versions={manifest.version: release}
            )
        },
        cache_root=tmp_path / "cache",
    )
    registry.build("riborf-mouse", source_paths=paths)
    return registry.open("riborf-mouse")


def test_native_versions_blocks_and_evidence_survive_target_placement(catalog):
    assert len(catalog.records) == 13
    for record in catalog.records:
        native = record.lifted_from
        mapping = record.reference_mapping
        assert record.liftover is None and native.genome_build == "GRCm38"
        assert record.genome_build == "GRCm39"
        assert (
            record.transcript_id != native.transcript_id
        )  # Real M25 -> Ensembl 107 version changes.
        assert record.protein_sequence == native.protein_sequence
        assert record.evidence == native.evidence
        assert all(s.target.reference == native.reference_context for s in record.summaries)
        assert mapping.source.identifier == native.orf_id == record.orf_id
        assert mapping.status == "unique" and mapping.sequence_validated
        assert mapping.validation["coding_sequence"]["exact"] is True
        assert mapping.validation["rna_sequence"]["exact"] is True
        assert mapping.provenance["source_row_id"] in native.orf_id
        assert _record(record.to_dict()) == record
    assert sorted(len(r.protein_sequence) for r in catalog.records)[:2] == [1, 4]
    assert len(catalog.query(min_protein_length=5)) == 11
    assert {r.biotype for r in catalog.records} == {
        "uORF",
        "uoORF",
        "iORF",
        "doORF",
        "lncORF",
        "extension",
        "truncation",
    }


def test_changed_or_missing_target_rna_keeps_native_models(catalog):
    unresolved = catalog.metadata["unresolved"]
    assert len(unresolved) == 4
    assert all(item["native_record"] for item in unresolved)
    assert all(item["reference_mapping"]["status"] == "unresolved" for item in unresolved)
    original = [_record(item["native_record"]) for item in unresolved]
    assert {r.transcript_id for r in original} == {
        "ENSMUST00000020909.3",
        "ENSMUST00000192542.1",
        "ENSMUST00000170111.1",
    }
    # The unchanged short CDS in the changed chr14 RNA is deliberately not used
    # to claim a uniquely preserved RNA/position across the duplicated region.
    assert not any(r.transcript_id.startswith("ENSMUST00000170111") for r in catalog.records)
    assert all(r.summaries and r.reference_context.annotation_version == "M25" for r in original)


def test_presence_remains_aggregate_native_ribo_seq(catalog):
    assert len(catalog.query(evidence=RiboSeq())) == 13
    assert catalog.query(evidence=MS()) == ()
    assert catalog.query(evidence=RiboSeq(min_studies=2)) == ()
    (summary,) = catalog.records[0].evidence.summaries
    assert (summary.metric, summary.value, summary.scope) == ("RibORF 2.0 translated ORF", 1, "orf")


def test_every_record_binds_to_its_target_transcript(catalog, tmp_path):
    from pyensembl import Genome
    from pyensembl.species import find_species_by_name

    gtf = tmp_path / "e107.gtf"
    gtf.write_bytes(gzip.decompress(PATHS["gtf"].read_bytes()))
    fastas = []
    for kind in ("cdna", "ncrna"):
        path = tmp_path / f"e107.{kind}.fa"
        path.write_bytes(gzip.decompress(PATHS[kind].read_bytes()))
        fastas.append(str(path))
    genome = Genome(
        "GRCm39",
        "ensembl",
        "107",
        gtf_path_or_url=str(gtf),
        transcript_fasta_paths_or_urls=fastas,
        cache_directory_path=str(tmp_path / "pyensembl"),
    )
    genome.species = find_species_by_name("mouse")
    genome.index()
    try:
        for record in catalog.records:
            bound = record.bind(genome.transcript_by_id(record.transcript_id.split(".")[0]))
            assert bound.protein_sequence == record.lifted_from.protein_sequence
        placed = extra_orfs(genome, None, catalog=catalog)
        assert placed.excluded == ()
        assert len(placed.unlifted) == 4
        selected = extra_orfs(genome, None, catalog=catalog, min_protein_length=5)
        assert len({orf.orf_id for orf in selected}) == 11 and selected.length_filtered == 2
        proteome = Proteome.build(
            genome, reference_cds=False, nc_orf_sources=catalog, min_protein_length=5
        )
        assert dict(proteome.layers[-1])["length_filtered"] == 2
    finally:
        genome.close()


def rewrite_fasta(path, sequences):
    text = "".join(f">{name}\n{sequence}\n" for name, sequence in sequences.items())
    path.write_bytes(gzip.compress(text.encode(), mtime=0))


@pytest.mark.parametrize("mutation", ["codon", "compensating_indels"])
def test_same_length_coding_changes_never_inherit_native_evidence(study, mutation):
    paths, manifest = study
    before, _, _ = reconstruct(paths, manifest)
    record = next(r for r in before if r.biotype == "truncation")
    interval = record.reference_mapping.provenance["source_rna_interval"]
    start, end = interval["start"], interval["end"]
    sequences = _fasta([paths["cdna"]])
    sequence = sequences[record.transcript_id]
    coding = sequence[start - 1 : end - 1]
    if mutation == "codon":
        replacement = "AAA" if coding[3:6] != "AAA" else "CCC"
        changed = coding[:3] + replacement + coding[6:]
    else:
        # A one-base insertion and a later deletion preserve the overall RNA
        # length and boundaries. Choose a case without an intermediate stop.
        changed = None
        for offset in range(8, 70):
            trial = coding[:6] + "C" + coding[6:offset] + coding[offset + 1 :]
            try:
                _, protein, _ = _coding(trial, 1, len(trial) + 1, trial[:3])
            except ValueError:
                continue
            if protein != record.protein_sequence:
                changed = trial
                break
        assert changed is not None
    assert len(changed) == len(coding) and changed[:3] == coding[:3] and changed[-3:] == coding[-3:]
    assert _coding(changed, 1, len(changed) + 1, changed[:3])[1] != record.protein_sequence
    sequences[record.transcript_id] = sequence[: start - 1] + changed + sequence[end - 1 :]
    rewrite_fasta(paths["cdna"], sequences)
    records, unresolved, _ = reconstruct(paths, repin(paths, manifest))
    assert record.orf_id not in {r.orf_id for r in records}
    retained = next(
        u for u in unresolved if u.get("native_record", {}).get("orf_id") == record.orf_id
    )
    assert retained["native_record"]["protein_sequence"] == record.protein_sequence
    assert retained["reference_mapping"]["targets"] == []


def test_duplicate_full_rna_loci_are_explicitly_ambiguous(study):
    paths, manifest = study
    before, _, _ = reconstruct(paths, manifest)
    record = next(r for r in before if r.biotype == "truncation")
    sequences = _fasta([paths["cdna"]])
    duplicate = "ENSMUST99999999999.7"
    sequences[duplicate] = sequences[record.transcript_id]
    rewrite_fasta(paths["cdna"], sequences)
    text = gzip.decompress(paths["gtf"].read_bytes()).decode()
    stable = record.transcript_id.split(".")[0]
    appended = []
    for line in text.splitlines():
        if f'transcript_id "{stable}"' in line:
            fields = line.split("\t")
            fields[3], fields[4] = str(int(fields[3]) + 1000000), str(int(fields[4]) + 1000000)
            fields[8] = fields[8].replace(stable, duplicate.split(".")[0])
            appended.append("\t".join(fields) + "\n")
    assert appended
    paths["gtf"].write_bytes(gzip.compress((text + "".join(appended)).encode(), mtime=0))
    records, unresolved, _ = reconstruct(paths, repin(paths, manifest))
    assert record.orf_id not in {r.orf_id for r in records}
    retained = next(
        u for u in unresolved if u.get("native_record", {}).get("orf_id") == record.orf_id
    )
    assert retained["reference_mapping"]["status"] == "ambiguous"
    assert len(retained["reference_mapping"]["targets"]) == 2
    assert len(retained["reference_mapping"]["provenance"]["candidate_positions"]) == 2


def test_conflicting_fasta_identity_is_not_silently_overwritten(study):
    paths, manifest = study
    sequences = _fasta([paths["cdna"]])
    name, sequence = next(iter(sequences.items()))
    text = gzip.decompress(paths["cdna"].read_bytes())
    paths["cdna"].write_bytes(
        gzip.compress(text + f">{name}\n{sequence[:-1]}A\n".encode(), mtime=0)
    )
    if sequence[-1] == "A":
        paths["cdna"].write_bytes(
            gzip.compress(text + f">{name}\n{sequence[:-1]}C\n".encode(), mtime=0)
        )
    with pytest.raises(ValueError, match="Conflicting FASTA"):
        reconstruct(paths, repin(paths, manifest))


def test_native_version_is_not_rebound_by_stable_id(study):
    paths, manifest = study
    sequences = _fasta([paths["source_rna"]])
    name = "ENSMUST00000000187.6"
    sequences[name.rsplit(".", 1)[0] + ".99"] = sequences.pop(name)
    rewrite_fasta(paths["source_rna"], sequences)
    records, unresolved, _ = reconstruct(paths, repin(paths, manifest))
    assert not any(r.lifted_from.transcript_id == name for r in records)
    item = next(u for u in unresolved if u.get("source_transcript") == name)
    assert item["native_record"] is None and "no source RNA" in item["reason"]


def test_native_bed_blocks_are_validated_before_placement(study):
    from openpyxl import load_workbook

    paths, manifest = study
    workbook = load_workbook(paths["table"])
    workbook["mouse"].cell(1, 11, "24,")  # Native interval is 21 nt.
    workbook.save(paths["table"])
    workbook.close()
    records, unresolved, _ = reconstruct(paths, repin(paths, manifest))
    assert len(records) == 12
    item = next(
        u
        for u in unresolved
        if u.get("source_row_id", "").startswith("ENSMUST00000000087.12:4:+|3|")
    )
    assert item["native_record"] is None and "BED" in item["reason"]


@pytest.mark.parametrize("mutation", ["endpoint", "position", "protein_hash"])
def test_portable_mapping_conflicts_fail_before_binding(catalog, tmp_path, mutation):
    data = catalog.to_dict()
    mapping = data["records"][0]["reference_mapping"]
    if mutation == "endpoint":
        mapping["source"]["identifier"] = "different-source-orf"
    elif mutation == "position":
        mapping["validation"]["position"]["target"]["coding_blocks"][0][0] += 1
    else:
        mapping["validation"]["sequence"]["source_sha256"] = "a" * 64
        mapping["validation"]["sequence"]["target_sha256"] = "a" * 64
    path = tmp_path / "changed.json"
    path.write_text(json.dumps(data))
    with pytest.raises(ValueError, match="reference mapping"):
        load_catalog(path)


def test_setup_reuses_the_shared_target_annotation_once(monkeypatch, tmp_path):
    from orfdata import install

    fetched = []
    monkeypatch.setattr(
        install,
        "_fetch",
        lambda url, sha256, size, staging, options: fetched.append(url) or "local",
    )
    registry = CatalogRegistry()
    manifests = [
        registry.info("ncorf-mammals-mouse"),
        registry.info("riborf-mouse"),
        registry.info("openprot-mouse"),
    ]
    shared = install._shared_sources(manifests, tmp_path, {})
    (gtf,) = (s for s in manifests[1].sources if s.asset_id == "gtf")
    assert shared == {gtf.sha256: "local"} and fetched == [gtf.url]


@pytest.mark.parametrize("version", ["1", "2"])
def test_superseded_unvalidated_mouse_builds_require_current_generation(version, tmp_path):
    from orfdata.sources.riborf import import_orfs

    registry = CatalogRegistry()
    assert registry.info("riborf-mouse").version == "3"
    with pytest.raises(ValueError, match="lack native sequence validation; build generation 3"):
        import_orfs({}, tmp_path / "old.gz", registry.info("riborf-mouse", version))
    assert not (tmp_path / "old.gz").exists()


def test_tsv_retains_typed_mapping_and_native_record(catalog, tmp_path, capsys):
    import csv
    import io

    from orfdata.cli import run
    from orfdata.reference_mapping import ReferenceMapping

    path = tmp_path / "portable.json.gz"
    catalog.save(path)
    run(["--catalog", str(path), "--format", "tsv"])
    rows = list(csv.DictReader(io.StringIO(capsys.readouterr().out), delimiter="\t"))
    assert len(rows) == 13
    by_id = {r.orf_id: r for r in catalog.records}
    for row in rows:
        record = by_id[row["orf_id"]]
        assert (
            ReferenceMapping.from_dict(json.loads(row["reference_mapping_json"]))
            == record.reference_mapping
        )
        assert _record(json.loads(row["lifted_from_json"])) == record.lifted_from


def test_registry_rejects_unvalidated_rebuild_before_acquisition(monkeypatch, tmp_path):
    import orfdata.catalogs as catalogs

    monkeypatch.setattr(
        catalogs,
        "build_catalog",
        lambda *a, **k: pytest.fail("Must reject before source acquisition"),
    )
    with pytest.raises(ValueError, match="lack native sequence validation"):
        CatalogRegistry(cache_root=tmp_path).build("riborf-mouse", "2")


def test_native_reference_manifest_contract_is_species_scoped(study):
    _, manifest = study
    assert CatalogManifest.from_dict(manifest.to_dict()) == manifest
    other = replace(manifest.source_references[0], taxon_id=9606, species="human")
    with pytest.raises(ValueError, match="catalogue species"):
        replace(manifest, source_references=(other,))
    # A source reference cannot substitute for the catalogue's target annotation.
    assert (
        manifest.genome_build == "GRCm39" and manifest.source_references[0].genome_build == "GRCm38"
    )


def test_merged_replay_preserves_each_mapping_without_duplicate_evidence(catalog, tmp_path):
    from orfdata import ORFCatalog, ORFRecord, merge_catalogs
    from orfdata.liftover import unlifted

    with pytest.raises(ValueError, match="merge one release of each source"):
        merge_catalogs([catalog, catalog])
    other = ORFRecord.from_model(
        replace(catalog.records[0].model, source="independent-fixture", orf_id="other-orf")
    )
    second = ORFCatalog(
        [other],
        dict(
            genome_build="GRCm39",
            taxon_id=10090,
            species="mus_musculus",
            coordinates="1-based-inclusive",
            dataset_version="fixture",
            sources=[dict(id="independent-fixture")],
        ),
    )
    merged = merge_catalogs([catalog, second])
    assert len(merged.records) == 13 and len(merged.evidence_summaries) == 13
    path = tmp_path / "merged.json.gz"
    merged.save(path)
    restored = load_catalog(path)
    assert restored.to_dict() == merged.to_dict()
    assert all(
        r.reference_mapping or any(m.reference_mapping for m in r.members) for r in restored.records
    )
    assert len({r.orf_id for r, _, _ in unlifted(restored.metadata)}) == 4


def test_declared_native_or_target_reference_cannot_relabel_pinned_inputs(study):
    paths, manifest = study
    for changed in (
        replace(manifest, references=(replace(manifest.references[0], annotation_version="108"),)),
        replace(
            manifest,
            source_references=(replace(manifest.source_references[0], genome_build="GRCm39"),),
        ),
    ):
        with pytest.raises(ValueError, match="reference identities"):
            reconstruct(paths, changed)


def test_unresolved_mapping_conflict_is_rejected_on_replay(catalog, tmp_path):
    data = catalog.to_dict()
    data["metadata"]["unresolved"][0]["reference_mapping"]["source"]["identifier"] = "other-orf"
    path = tmp_path / "wrong-unresolved.json"
    path.write_text(json.dumps(data))
    with pytest.raises(ValueError, match="Unresolved reference mapping"):
        load_catalog(path)


def test_unresolved_reference_must_belong_to_the_managed_manifest(catalog, tmp_path):
    data = catalog.to_dict()
    data["metadata"]["unresolved"][0]["reference_mapping"]["target_reference"][
        "annotation_version"
    ] = "108"
    path = tmp_path / "wrong-target-release.json"
    path.write_text(json.dumps(data))
    changed = load_catalog(path)
    manifest = CatalogManifest.from_dict(changed.metadata["catalog_manifest"])
    with pytest.raises(ValueError, match="Unresolved placement reference"):
        manifest.validate_catalog(changed)


def test_later_chain_lift_keeps_the_original_typed_rna_mapping(catalog, tmp_path):
    from orfdata import ChainFile, ORFCatalog, lift_catalog

    record = next(r for r in catalog.records if len(r.coding_blocks) == 1)
    first, last = record.coding_blocks[0]
    chain_path = tmp_path / "synthetic.chain"
    chain_path.write_text(
        f"chain 1 {record.contig} 200000000 + {first - 1} {last} "
        f"{record.contig} 200000000 + {first - 1} {last} 1\n{last - first + 1}\n"
    )
    chain = ChainFile(chain_path, source_build="GRCm39", target_build="GRCm38")
    single = ORFCatalog(
        [record],
        dict(
            genome_build="GRCm39",
            taxon_id=10090,
            species="mus_musculus",
            coordinates="1-based-inclusive",
            dataset_version="chain-fixture",
        ),
    )
    lifted = lift_catalog(single, chain).records[0]
    assert lifted.reference_mapping is None and lifted.liftover == chain.liftover
    assert lifted.lifted_from.reference_mapping == record.reference_mapping
    assert lifted.lifted_from.lifted_from == record.lifted_from
    assert _record(lifted.to_dict()) == lifted
