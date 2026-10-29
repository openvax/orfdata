"""Independently audit every native mouse RibORF and target placement.

Uses Biopython translation and direct base-by-base GTF coordinate comparison;
requires the development dependencies, not additional reference indexes.
"""

import argparse
import gzip
import hashlib
import itertools
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from Bio import SeqIO
from Bio.Seq import Seq
from openpyxl import load_workbook


def rnas(paths):
    result = {}
    for path in paths:
        with gzip.open(path, "rt") as stream:
            for record in SeqIO.parse(stream, "fasta"):
                name = record.id.split("|")[0]
                sequence = str(record.seq).upper()
                assert name not in result or result[name] == sequence
                result[name] = sequence
    return result


def exons(path, wanted):
    result = defaultdict(list)
    with gzip.open(path, "rt") as stream:
        for line in stream:
            if line.startswith("#"):
                continue
            fields = line.rstrip("\n").split("\t")
            if fields[2] != "exon":
                continue
            attributes = dict(re.findall(r'(\w+) "([^"]*)"', fields[8]))
            name = attributes["transcript_id"]
            if "." not in name and attributes.get("transcript_version"):
                name += "." + attributes["transcript_version"]
            if name in wanted:
                result[name].append(
                    (fields[0].removeprefix("chr"), fields[6], int(fields[3]), int(fields[4]))
                )
    return result


def check_path(record, annotations, start, end):
    entries = annotations[record["transcript_id"]]
    assert {(c, strand) for c, strand, _, _ in entries} == {(record["contig"], record["strand"])}
    reverse = record["strand"] == "-"
    spans = sorted({(a, b) for _, _, a, b in entries}, reverse=reverse)
    coordinates = itertools.chain.from_iterable(
        range(b, a - 1, -1) if reverse else range(a, b + 1) for a, b in spans
    )
    actual = list(itertools.islice(coordinates, start - 1, end - 1))
    blocks = sorted(record["coding_blocks"], reverse=reverse)
    expected = list(
        itertools.chain.from_iterable(
            range(b, a - 1, -1) if reverse else range(a, b + 1) for a, b in blocks
        )
    )
    assert actual == expected


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("catalog", type=Path)
    parser.add_argument(
        "source_paths", type=Path, help="JSON mapping asset IDs to local source paths"
    )
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    paths = json.loads(args.source_paths.read_text())
    with gzip.open(args.catalog, "rt") as stream:
        catalog = json.load(stream)
    manifest = catalog["metadata"]["catalog_manifest"]
    for asset in manifest["sources"]:
        with open(paths[asset["asset_id"]], "rb") as stream:
            assert hashlib.file_digest(stream, "sha256").hexdigest() == asset["sha256"]
    source = rnas([paths["source_rna"]])
    target = rnas([paths["cdna"], paths["ncrna"]])
    by_sequence = defaultdict(set)
    for name, sequence in target.items():
        by_sequence[hashlib.sha256(sequence.encode()).digest()].add(name)
    workbook = load_workbook(paths["table"], read_only=True)
    rows = {
        row[3]: list(row)
        for row in workbook["mouse"].iter_rows(values_only=True)
        if row and row[3] and row[3].split("|")[3] not in ("canonical", "seqerror", "readthrough")
    }
    workbook.close()
    all_records = [
        (record["lifted_from"], record["reference_mapping"], record)
        for record in catalog["records"]
    ]
    all_records += [
        (item["native_record"], item["reference_mapping"], None)
        for item in catalog["metadata"]["unresolved"]
    ]
    assert len(all_records) == len(rows) == 54128
    source_annotation = exons(paths["source_gtf"], {r[0]["transcript_id"] for r in all_records})
    target_annotation = exons(paths["gtf"], {r[2]["transcript_id"] for r in all_records if r[2]})
    seen, statuses = set(), Counter()
    for native, mapping, placed in all_records:
        provenance = mapping["provenance"]
        row_id = provenance["source_row_id"]
        assert row_id not in seen
        seen.add(row_id)
        row = rows[row_id]
        assert provenance["source_row"] == row
        transcript, _, interval, _, codon = row_id.split("|")
        assert native["transcript_id"] == transcript.split(":")[0]
        length, start, end = map(int, interval.split(":"))
        sequence = source[native["transcript_id"]]
        assert len(sequence) == length
        coding = sequence[start - 1 : end - 1]
        assert provenance["source_coding_sequence"] == coding and coding[:3] == codon
        translated = str(Seq(coding).translate(table=1))
        assert "M" + translated[1:] == native["protein_sequence"] + "*"
        assert "*" not in translated[:-1]
        sizes = [int(n) for n in row[10].rstrip(",").split(",")]
        starts = [int(n) for n in row[11].rstrip(",").split(",")]
        blocks = [
            [row[1] + offset + 1, row[1] + offset + size] for offset, size in zip(starts, sizes)
        ]
        assert blocks == native["coding_blocks"]
        check_path(native, source_annotation, start, end)
        candidates = by_sequence[hashlib.sha256(sequence.encode()).digest()]
        assert all(target[name] == sequence for name in candidates)
        if placed is not None:
            assert candidates == {placed["transcript_id"]}
            assert target[placed["transcript_id"]][start - 1 : end - 1] == coding
            assert placed["protein_sequence"] == native["protein_sequence"]
            assert placed["summaries"] == native["summaries"]
            check_path(placed, target_annotation, start, end)
            assert mapping["status"] == "unique"
        else:
            assert mapping["status"] == ("ambiguous" if len(candidates) > 1 else "unresolved")
            assert len(candidates) != 1
        statuses[mapping["status"]] += 1
    assert seen == set(rows)
    receipt = dict(
        source_rows_checked=len(seen),
        statuses=dict(statuses),
        translation="All native CDSs independently translated with Biopython",
        geometry="All native and accepted target coding bases compared with GTF exon paths",
        target_scope="All pinned cDNA/ncRNA sequences searched; exact bytes compared",
        catalog_sha256=hashlib.sha256(args.catalog.read_bytes()).hexdigest(),
    )
    args.output.write_text(json.dumps(receipt, indent=2) + "\n")
    print(json.dumps(receipt, indent=2))


if __name__ == "__main__":
    main()
