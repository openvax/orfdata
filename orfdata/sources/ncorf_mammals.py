"""Ribo-seq ORFs of Chang et al., "Evolutionary remodeling of non-canonical ORF translation in mammals".

The source called ORFs with PRICE, RiboCode and Ribo-TISH on Ensembl 107
annotations. An ORF was kept when at least two callers found it in a library,
in at least two libraries; starts are ATG, CTG, GTG or TTG. Its BED file
garbles the block sizes of multi-exon ORFs, so coding blocks come from each
ORF's representative transcript and transcript coordinates (Table S3), mapped
through the pinned Ensembl 107 GTF, and are checked against the ORF's
genomic start and stop, the BED span and the protein length. ORFs that fail,
such as one without a stop codon, are listed in ``metadata["unresolved"]``.

The source writes GTG-initiated proteins with valine; a near-cognate start is
read by the initiator methionine, so the protein here begins with M and the
record declares its start codon. The library and caller counts are aggregate
Ribo-seq support for the ORF: they name no studies.
"""

import gzip
import hashlib
import json
import re
from collections import Counter

from ..catalog import ORFCatalog, ORFEvidence, ORFRecord
from ..evidence_summary import EvidenceSummary
from ..model import ORFModel
from ..reference import ScopedTarget
from ..sequence import genetic_code

LIBRARIES = ("Ribo-seq libraries in which at least two of PRICE, RiboCode and Ribo-TISH called the ORF "
             "(Table S3 num_libraries); the libraries are listed in Table S1, not per ORF.")
CALLERS = "ORF callers, of PRICE, RiboCode and Ribo-TISH, that called the ORF (Table S3 num_methods)."
_ATTRIBUTE = re.compile(r'(\S+) "([^"]*)"')


def _fasta(path):
    sequences, name = {}, None
    with open(path) as handle:
        for line in handle:
            line = line.strip()
            if line.startswith(">"):
                name = line[1:].split()[0]
                sequences[name] = []
            elif name is not None:
                sequences[name].append(line)
    return {k: "".join(v) for k, v in sequences.items()}


def _table(path, sheet, species):
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ImportError("Install orfdata[build-data] to build this catalogue from its spreadsheet") from error

    with open(path, "rb") as handle:  # Cached files have no extension for openpyxl to go by.
        workbook = load_workbook(handle, read_only=True)
        rows = workbook[sheet].iter_rows(values_only=True)
        header = next(rows)
        found = [dict(zip(header, row)) for row in rows if row and row[0] == species]
        workbook.close()
    return found


def _transcripts(path, wanted):
    """Exons in transcript order, with contig, strand, version and gene, for the wanted transcripts."""
    found = {}
    with gzip.open(path, "rt") as handle:
        for line in handle:
            if line.startswith("#"):
                continue
            fields = line.rstrip("\n").split("\t")
            if len(fields) < 9 or fields[2] != "exon":
                continue
            attributes = dict(_ATTRIBUTE.findall(fields[8]))
            transcript = attributes.get("transcript_id")
            if transcript not in wanted:
                continue
            entry = found.setdefault(transcript, dict(
                contig=fields[0], strand=fields[6], version=attributes.get("transcript_version"),
                gene_id=attributes["gene_id"], gene_name=attributes.get("gene_name", ""), exons=[]))
            entry["exons"].append((int(fields[3]), int(fields[4])))
    for entry in found.values():
        entry["exons"].sort(reverse=entry["strand"] == "-")
    return found


def _blocks(entry, start, end):
    """1-based inclusive genomic blocks, ascending, of transcript positions ``start``..``end``."""
    blocks, offset = [], 0
    for a, b in entry["exons"]:
        length = b - a + 1
        lo, hi = max(start, offset + 1), min(end, offset + length)
        if lo <= hi:
            if entry["strand"] == "+":
                blocks.append((a + lo - offset - 1, a + hi - offset - 1))
            else:
                blocks.append((b - (hi - offset - 1), b - (lo - offset - 1)))
        offset += length
    return tuple(sorted(blocks)) if offset >= end else None


def import_orfs(inputs, output, manifest):
    """Write the catalogue ``manifest`` describes from the source's table, BED, proteins and the Ensembl GTF."""
    processing = manifest.processing
    reference = manifest.references[0]
    table = next(s for s in manifest.sources if s.asset_id == "table")
    rows = _table(inputs["table"], processing["sheet"], processing["species_label"])
    proteins = _fasta(inputs["proteins"])
    spans = {}
    with open(inputs["bed"]) as handle:
        for line in handle:
            fields = line.rstrip("\n").split("\t")
            if len(fields) >= 6:
                spans[fields[3]] = (fields[0], int(fields[1]) + 1, int(fields[2]), fields[5])
    transcripts = _transcripts(inputs["gtf"], {r["representative_transcript"] for r in rows})

    records, unresolved = [], []
    for row in rows:
        orf_id, transcript = row["orf_id"], row["representative_transcript"]

        def skip(reason):
            unresolved.append(dict(orf_id=orf_id, transcript=transcript, reason=reason))

        entry, protein = transcripts.get(transcript), proteins.get(orf_id)
        if entry is None:
            skip(f"Transcript absent from Ensembl {reference.annotation_version}")
            continue
        if protein is None or orf_id not in spans:
            skip("ORF missing from the source's protein FASTA or BED")
            continue
        blocks = _blocks(entry, int(row["transcript_start"]), int(row["transcript_end"]))
        strand = "+" if row["strand"] == "+" else "-"
        first, last = (int(row["genomic_start"]), int(row["genomic_end"]))
        if (blocks is None or entry["strand"] != strand or str(row["chromosome"]) != entry["contig"]
                or (blocks[0][0], blocks[-1][1]) != ((first, last) if strand == "+" else (last, first))
                or spans[orf_id] != (entry["contig"], blocks[0][0], blocks[-1][1], strand)):
            skip("Transcript coordinates disagree with the ORF's genomic start, stop or BED span")
            continue
        if sum(b - a + 1 for a, b in blocks) != 3 * (len(protein) + 1):
            skip("Coding length is not the protein plus a stop codon")
            continue
        start_codon = row["start_codon"]
        if start_codon != "ATG":
            protein = "M" + protein[1:]  # A near-cognate start is read by the initiator methionine.
        model = ORFModel(orf_id, f"{transcript}.{entry['version']}", entry["gene_id"], entry["gene_name"],
                         entry["contig"], strand, blocks, protein, row["orf_type"], reference.genome_build,
                         reference.taxon_id, reference.species, source=processing["source"],
                         reference_context=reference,
                         start_codon=None if start_codon in genetic_code(1).start_codons else start_codon)
        target = ScopedTarget(reference, "orf", orf_id, processing["source"])
        summaries = tuple(
            EvidenceSummary(processing["source"], processing["release"], target, _sha(protein), "ribosome", metric,
                            int(value), definition, "orf", table.url, table.sha256)
            for metric, value, definition in (("Ribo-seq libraries", row["num_libraries"], LIBRARIES),
                                              ("ORF callers", row["num_methods"], CALLERS)))
        records.append(ORFRecord.from_model(model, ORFEvidence(summaries=summaries)))

    counts = dict(source_rows=len(rows), resolved_rows=len(records), unresolved_rows=len(unresolved))
    for key, value in counts.items():
        if manifest.coverage.get(key) != value:
            raise ValueError(f"Source {key} is {value}, but the manifest declares {manifest.coverage.get(key)}")
    metadata = dict(
        genome_build=reference.genome_build, taxon_id=reference.taxon_id, species=reference.species,
        coordinates="1-based-inclusive", dataset_version=processing["release"],
        sources=[dict(id=processing["source"], release=processing["release"], url=table.url,
                      scope=processing["scope"])],
        unresolved=unresolved, unresolved_reasons=dict(Counter(u["reason"] for u in unresolved)),
        catalog_manifest=manifest.to_dict())
    catalog = ORFCatalog(records, metadata)
    raw = json.dumps(catalog.to_dict(), sort_keys=True, separators=(",", ":"), allow_nan=False).encode()
    with open(output, "wb") as handle, gzip.GzipFile(filename="", mode="wb", fileobj=handle, mtime=0) as compressed:
        compressed.write(raw)


def _sha(protein):
    return hashlib.sha256(protein.encode("ascii")).hexdigest()
